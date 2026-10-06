from __future__ import annotations

import os
import unittest
import uuid
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from backend.api.server import app


class WorkbookModelApiTests(unittest.TestCase):
    def setUp(self) -> None:
        runtime_dir = Path(__file__).resolve().parent.parent / "runtime"
        runtime_dir.mkdir(parents=True, exist_ok=True)
        self.workbook_path = runtime_dir / f"memory-model-{uuid.uuid4().hex}.xlsx"

        workbook = Workbook()
        hbm = workbook.active
        hbm.title = "DRAM Demand"
        hbm["B2"] = "DRAM Memory Demand based on GW Provisioning"
        hbm["D3"] = 2025
        hbm["E3"] = 2026
        hbm["B5"] = "Arete"
        hbm["D5"] = 11.7
        hbm["E5"] = 22.6
        hbm["B6"] = "% yoy"
        hbm["E6"] = 0.931623931
        hbm["E6"].number_format = "0.0%"
        hbm["I144"] = 0.317397195
        hbm["I144"].number_format = "0.00%"

        capex = workbook.create_sheet("Hyperscaler Capex GW Triangulat")
        for column, year in zip(range(4, 10), range(2025, 2031), strict=True):
            capex.cell(2, column, year)
        capex["B26"] = "Datacentre Capex USD Mn"
        capex["B27"] = "AMZN US Equity"
        capex["D27"] = -34194.0064395215
        capex["I27"] = -285329.8
        capex["B28"] = "GOOGL US Equity"
        capex["D28"] = -56390.4299501731
        capex["I28"] = -342793.9333333333
        capex["B34"] = "Hyperscaler Capex total"
        capex["D34"] = -90584.4363896946
        capex["B35"] = "% yoy"
        capex["E35"] = 1.405165727
        capex["E35"].number_format = "0.0%"
        workbook.save(self.workbook_path)

        self.env = patch.dict(
            os.environ,
            {"MEMORY_MODEL_WORKBOOK_PATH": str(self.workbook_path)},
            clear=False,
        )
        self.env.start()
        self.client = TestClient(app, raise_server_exceptions=False)

    def tearDown(self) -> None:
        self.client.close()
        self.env.stop()
        self.workbook_path.unlink(missing_ok=True)

    def test_hbm_table_reads_first_sheet_and_formats_percentages(self) -> None:
        response = self.client.get("/api/model/workbook")

        self.assertEqual(200, response.status_code)
        self.assertEqual("application/json", response.headers.get("content-type"))
        payload = response.json()
        self.assertEqual(self.workbook_path.name, payload["source"]["filename"])
        self.assertEqual("DRAM Demand", payload["hbm"]["sheet"])
        self.assertEqual("B2:I144", payload["hbm"]["range"])
        self.assertEqual(143, len(payload["hbm"]["rows"]))
        self.assertEqual(
            "DRAM Memory Demand based on GW Provisioning",
            payload["hbm"]["rows"][0]["cells"][0]["display"],
        )
        self.assertEqual("93.16%", payload["hbm"]["rows"][4]["cells"][3]["display"])
        self.assertEqual("31.74%", payload["hbm"]["rows"][142]["cells"][7]["display"])

    def test_capex_payload_preserves_negative_values_for_grouped_bars(self) -> None:
        response = self.client.get("/api/model/workbook")

        self.assertEqual(200, response.status_code)
        self.assertEqual("application/json", response.headers.get("content-type"))
        capex = response.json()["capex"]
        self.assertEqual("Hyperscaler Capex GW Triangulat", capex["sheet"])
        self.assertEqual("B26:I35", capex["range"])
        self.assertEqual([2025, 2026, 2027, 2028, 2029, 2030], capex["years"])
        self.assertEqual(
            {
                "name": "AMZN US Equity",
                "values": [-34194.0064395215, None, None, None, None, -285329.8],
            },
            capex["series"][0],
        )
        self.assertEqual(-56390.4299501731, capex["series"][1]["values"][0])
        self.assertEqual("140.52%", capex["table"]["rows"][9]["cells"][3]["display"])

    def test_changed_workbook_mtime_reloads_cell_values(self) -> None:
        first = self.client.get("/api/model/workbook").json()
        old_mtime = self.workbook_path.stat().st_mtime_ns
        workbook = load_workbook(self.workbook_path)
        workbook["DRAM Demand"]["D5"] = 99.25
        workbook.save(self.workbook_path)
        os.utime(self.workbook_path, ns=(old_mtime + 2_000_000_000, old_mtime + 2_000_000_000))

        second = self.client.get("/api/model/workbook").json()

        self.assertEqual(11.7, first["hbm"]["rows"][3]["cells"][2]["value"])
        self.assertEqual(99.25, second["hbm"]["rows"][3]["cells"][2]["value"])
        self.assertNotEqual(first["source"]["modified_at"], second["source"]["modified_at"])

    def test_missing_workbook_returns_safe_explicit_error(self) -> None:
        self.workbook_path.unlink()

        response = self.client.get("/api/model/workbook")

        self.assertEqual(503, response.status_code)
        self.assertEqual("Memory model workbook not found.", response.json()["detail"])
        self.assertNotIn(str(self.workbook_path.parent), response.text)

    def test_corrupt_workbook_returns_explicit_read_error(self) -> None:
        self.workbook_path.write_bytes(b"not-an-xlsx")

        response = self.client.get("/api/model/workbook")

        self.assertEqual(503, response.status_code)
        self.assertEqual("Memory model workbook could not be read.", response.json()["detail"])


if __name__ == "__main__":
    unittest.main()
