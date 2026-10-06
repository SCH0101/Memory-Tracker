from __future__ import annotations

import os
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.utils import get_column_letter


DEFAULT_WORKBOOK_PATH = Path(
    r"C:\Users\mchso\OneDrive - Jebsen & Co Ltd\JCAPSharePoint - Documents"
    r"\Equity Team\Stocks\Coverage\000660 KS SK Hynix"
    r"\20260702 Memory Wafer per 1GW datacentre v2.xlsx"
)
CAPEX_SHEET = "Hyperscaler Capex GW Triangulat"

_cache_key: tuple[str, int] | None = None
_cache_payload: dict[str, Any] | None = None


def workbook_path() -> Path:
    configured = os.getenv("MEMORY_MODEL_WORKBOOK_PATH", "").strip()
    return Path(configured) if configured else DEFAULT_WORKBOOK_PATH


def _is_percentage(cell: Cell) -> bool:
    number_format = getattr(cell, "number_format", "")
    return isinstance(number_format, str) and "%" in number_format


def _display_value(cell: Cell) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (int, float)):
        if _is_percentage(cell):
            return f"{value * 100:.2f}%"
        if float(value).is_integer():
            return f"{int(value):,}"
        return f"{value:,.2f}"
    return str(value)


def _json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def _cell_payload(cell: Cell, address: str) -> dict[str, Any]:
    return {
        "address": address,
        "value": _json_value(cell.value),
        "display": _display_value(cell),
        "is_percentage": _is_percentage(cell),
    }


def _table_payload(sheet, min_row: int, max_row: int) -> dict[str, Any]:
    min_col, max_col = 2, 9
    return {
        "sheet": sheet.title,
        "range": f"B{min_row}:I{max_row}",
        "columns": [get_column_letter(column) for column in range(min_col, max_col + 1)],
        "rows": [
            {
                "row": row,
                "cells": [
                    _cell_payload(sheet.cell(row, column), f"{get_column_letter(column)}{row}")
                    for column in range(min_col, max_col + 1)
                ],
            }
            for row in range(min_row, max_row + 1)
        ],
    }


def _capex_payload(sheet) -> dict[str, Any]:
    table = _table_payload(sheet, 26, 35)
    years = [_json_value(sheet.cell(2, column).value) for column in range(4, 10)]
    series = []
    for row in range(27, 34):
        name = sheet.cell(row, 2).value
        if not name:
            continue
        series.append(
            {
                "name": str(name),
                "values": [_json_value(sheet.cell(row, column).value) for column in range(4, 10)],
            }
        )
    return {**table, "table": table, "years": years, "series": series}


def load_workbook_model() -> dict[str, Any]:
    global _cache_key, _cache_payload

    path = workbook_path()
    stat = path.stat()
    key = (str(path.resolve()), stat.st_mtime_ns)
    if key == _cache_key and _cache_payload is not None:
        return _cache_payload

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        hbm_sheet = workbook[workbook.sheetnames[0]]
        if CAPEX_SHEET not in workbook.sheetnames:
            raise ValueError(f"Workbook sheet not found: {CAPEX_SHEET}")
        payload = {
            "source": {
                "filename": path.name,
                "modified_at": datetime.fromtimestamp(stat.st_mtime).astimezone().isoformat(timespec="seconds"),
            },
            "hbm": _table_payload(hbm_sheet, 2, 144),
            "capex": _capex_payload(workbook[CAPEX_SHEET]),
        }
    finally:
        workbook.close()

    _cache_key = key
    _cache_payload = payload
    return payload
